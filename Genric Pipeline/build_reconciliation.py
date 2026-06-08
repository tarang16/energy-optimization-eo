"""
Master KPI reconciliation — network-agnostic.

Canonical KPI identity = (Plant, Network, Element, Instance, KPI Name).
For every KPI in the system-KPI export, reconcile across all sources:

  MAPPED?     -> do its PI tags resolve in EO_PI_Database
  FORMULA?    -> is there a formula in the SQL registry (per-unit duty/ENPI) or the
                 inferred sheet (per-metric: SEC / efficiency / CO2 / fuel input)
  CALCULATED? -> does the package produce a value

Outputs Data/source/KPI_Reconciliation.xlsx + a coverage summary, so you can see at a
glance how many SEU KPIs are mapped and which actually calculate.
"""
from __future__ import annotations
import csv, re, openpyxl
from collections import defaultdict
from string import ascii_uppercase
from pathlib import Path
from openpyxl import Workbook
from seu_engine import SEUEngine, REF
import kpi_map as KM

ROOT = Path(r"C:\Users\tnigam\Desktop\Python EO\Genric Pipeline\Data\source")
EXPORT = Path(r"C:\Users\tnigam\Downloads\system_kpis_all_2026-06-04 (2).xlsx")

# UI element -> compatible SQL seu_category (prevents cross-equipment mismatches)
ELEMENT_CATS = {
    "Furnace": {"Furnace"}, "Fuel Fired Boiler": {"Boiler"},
    "Steam Exchanger": {"Exchanger", "Reboiler"},
    "Backpressure Turbine": {"Compressor"},
    "Extraction-Condensing Turbine": {"Compressor"},
    "Feed Preheater": {"Exchanger", "Heater"},
}


def main():
    se = SEUEngine(); eng = se.eng
    pdb = openpyxl.load_workbook(ROOT / "PI_Database_EO.xlsx", data_only=True, read_only=True)["master_pi"]
    pi2name = {str(r[1]).strip(): str(r[0]).strip() for r in list(pdb.iter_rows(values_only=True))[1:] if r[1]}

    # SQL SEU leaf-sets (which PI-DB tags each unit's duty depends on) + duty value
    seu_info = {}
    for r in se.seus:
        expr = r.get("actual_duty_expression_gjph", "")
        refs = {se._alias(x) for x in REF.findall(str(expr))}
        leaves = set()
        for ref in refs:
            leaves |= eng.leaves(ref) if ref in eng.formulas else {ref}
        seu_info[r["seu_name"]] = {
            "leaves": leaves, "display": r["seu_display_name"],
            "cat": r["seu_category"], "duty": se.eval_expr(expr)}

    # ── semantic matcher: input-fingerprint (unit) + concept (metric) ──────────
    def toks(s): return {x for x in re.split(r"[^a-z0-9]+", str(s).lower()) if x}
    CONCEPT = {
        "isentropic efficiency": {"efficiency", "eff", "isentropic", "percent"},
        "efficiency": {"efficiency", "eff", "percent"},
        "sec": {"sp_en", "spec_en", "sec"}, "fuel input": {"duty", "energy", "fuel"},
        "co2": {"co2"}, "delivered power": {"power", "kw", "shaft"},
        "stage power": {"power", "stage"}, "stage efficiency": {"efficiency", "eff", "stage"},
        "steam rate": {"rate", "steam"}, "duty": {"duty", "energy"}, "power": {"power", "kw", "shaft"},
    }
    def concept_tokens(kpi):
        low = kpi.lower(); out = set()
        for k, v in CONCEPT.items():
            if all(w in low for w in k.split()): out |= v
        return out
    tag_tokens = {t: toks(t) for t in eng.formulas}
    def semantic(I, kpi):
        ct = concept_tokens(kpi)
        if not ct or not I:
            return None, None
        best = None
        for t in eng.formulas:
            L = eng.leaves(t)
            inter = I & L
            if not inter or not (ct & tag_tokens[t]):
                continue
            v = eng.eval(t)
            if v is None:
                continue
            score = (len(ct & tag_tokens[t]), len(inter) / len(L), len(inter) / len(I))
            if best is None or score > best[0]:
                best = (score, t, v)
        return (best[2], f"inferred:{best[1]}") if best else (None, None)

    # curated per-metric values (Furnace/Boiler) from kpi_map
    def curated_value(element, inst_letter, kpi):
        spec = KM.CURATED.get(element)
        if not spec or inst_letter not in spec["instances"]:
            return None, None
        idx = spec["instances"].index(inst_letter) + 1
        for lbl, vs in spec["kpis"].items():
            if lbl.split(" (")[0].lower() in kpi.lower():
                v, src = KM.value_for(eng, vs, idx)
                return v, src
        return None, None

    # read export, group by canonical identity
    ws = openpyxl.load_workbook(EXPORT, data_only=True, read_only=True)["SEU KPIs"]
    rows = list(ws.iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(rows) if r and r[0] == "Plant")
    h = {v: i for i, v in enumerate(rows[hi]) if v}; ci = lambda k: h[k]
    kpis = defaultdict(lambda: {"tags": set()})
    for r in rows[hi + 1:]:
        if not r or not r[ci("KPI Name")] or not r[ci("Element")]:
            continue
        key = (str(r[ci("Plant")]).strip(), str(r[ci("Network")]).strip(),
               str(r[ci("Element")]).strip(), str(r[ci("Instance")]).strip(),
               str(r[ci("KPI Name")]).strip())
        if r[ci("PI Tag")]:
            for tok in str(r[ci("PI Tag")]).split():
                kpis[key]["tags"].add(tok.strip())

    out = Workbook(); o = out.active; o.title = "reconciliation"
    o.append(["Plant", "Network", "Element", "Instance", "KPI Name",
              "#Tags", "TagsResolve", "MatchedUnit(SQL)", "FormulaSource", "Value", "Status"])

    summary = defaultdict(lambda: {"n": 0, "mapped": 0, "calc": 0})
    for (plant, net, elem, inst, kpi), d in sorted(kpis.items()):
        tags = d["tags"]
        ntags = len(tags)
        resolved = {pi2name[t] for t in tags if t in pi2name}
        nres = len(resolved)
        mapped = ntags > 0 and nres > 0
        # match SQL unit by input-overlap, restricted to compatible categories
        compat = ELEMENT_CATS.get(elem, None)
        cands = [(len(resolved & v["leaves"]), k) for k, v in seu_info.items()
                 if compat is None or v["cat"] in compat]
        best = max(cands, default=(0, None))
        unit = best[1] if best[0] > 0 else None
        # resolve value: curated metric first, else SQL duty ONLY for same-dimension (GJ/h) KPIs
        inst_letter = inst.strip().split()[-1]
        inst_letter = inst_letter if len(inst_letter) == 1 and inst_letter.isalpha() else ""
        val, src = curated_value(elem, inst_letter, kpi)
        is_gjph = ("gj/h" in kpi.lower()) or ("fuel input" in kpi.lower()) or ("duty" in kpi.lower())
        if val is None and unit and seu_info[unit]["duty"] is not None and is_gjph:
            val, src = seu_info[unit]["duty"], f"SQL:{unit}.actual_duty"
        if val is None:                       # 3rd source: semantic inferred match
            val, src = semantic(resolved, kpi)
        calc = val is not None
        status = "CALCULATED" if calc else ("MAPPED_NO_FORMULA" if mapped else "NOT_MAPPED")
        o.append([plant, net, elem, inst, kpi, ntags, f"{nres}/{ntags}",
                  unit, src or "", round(val, 3) if isinstance(val, float) else val, status])
        s = summary[elem]; s["n"] += 1
        if mapped: s["mapped"] += 1
        if calc: s["calc"] += 1

    out.save(ROOT / "KPI_Reconciliation.xlsx")
    print(f"{'Element':<32}{'KPIs':>6}{'mapped':>8}{'calculated':>12}")
    tn = tm = tc = 0
    for e, s in sorted(summary.items()):
        print(f"{e:<32}{s['n']:>6}{s['mapped']:>8}{s['calc']:>12}")
        tn += s["n"]; tm += s["mapped"]; tc += s["calc"]
    print(f"{'TOTAL':<32}{tn:>6}{tm:>8}{tc:>12}")
    print(f"\nWrote {ROOT / 'KPI_Reconciliation.xlsx'}")


if __name__ == "__main__":
    main()
