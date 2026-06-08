"""
config/loader.py
----------------
Loads KPI metadata from the bundled kpi_config.json or from an Excel file
(KPI_Configuration_*.xlsx).

KPIConfig provides per-KPI metadata:
    kpi_id, kpi_name, category, uom, design, min, max, default
"""
from __future__ import annotations

import json
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

_DATA_DIR = Path(__file__).parent.parent / "data"


@dataclass
class KPIConfig:
    kpi_id:      str
    kpi_name:    str
    category:    str
    uom:         str        = ""
    design:      Any        = None
    min_val:     Any        = None
    max_val:     Any        = None
    default_val: Any        = None

    def within_limits(self, value: float) -> bool:
        """Return True if value is between min and max (if defined)."""
        try:
            if self.min_val is not None and value < float(self.min_val):
                return False
            if self.max_val is not None and value > float(self.max_val):
                return False
            return True
        except (TypeError, ValueError):
            return True


def load_kpi_config(
    source: str | Path | None = None,
) -> dict[str, KPIConfig]:
    """
    Load KPI config from:
      - bundled kpi_config.json  (default, no source)
      - a custom JSON file
      - an Excel file (.xlsx)    — reads 'KPI Configuration' sheet

    Returns
    -------
    dict mapping kpi_id -> KPIConfig
    """
    if source is None:
        source = _DATA_DIR / "kpi_config.json"

    source = Path(source)
    if source.suffix.lower() in (".xlsx", ".xls"):
        return _from_excel(source)
    return _from_json(source)


def _from_json(path: Path) -> dict[str, KPIConfig]:
    with open(path) as f:
        rows = json.load(f)
    return {r["kpi_id"]: KPIConfig(
        kpi_id      = r["kpi_id"],
        kpi_name    = r.get("kpi_name", r["kpi_id"]),
        category    = r.get("category", ""),
        uom         = r.get("uom", ""),
        design      = r.get("design"),
        min_val     = r.get("min"),
        max_val     = r.get("max"),
        default_val = r.get("default"),
    ) for r in rows}


def _from_excel(path: Path) -> dict[str, KPIConfig]:
    try:
        import openpyxl
    except ImportError as e:
        raise ImportError("openpyxl required to load Excel KPI config") from e

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["KPI Configuration"]
    headers = [str(c.value).strip() if c.value else "" for c in next(ws.iter_rows(min_row=1, max_row=1))]
    col = {h: i for i, h in enumerate(headers)}

    configs = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        kpi_id = row[col.get("KPI ID", 1)]
        if not kpi_id:
            continue
        configs[str(kpi_id)] = KPIConfig(
            kpi_id      = str(kpi_id),
            kpi_name    = str(row[col.get("KPI Name", 2)] or kpi_id),
            category    = str(row[col.get("Category", 0)] or ""),
            uom         = str(row[col.get("UOM", 5)] or ""),
            design      = row[col.get("Design Data", 6)],
            min_val     = row[col.get("Min", 7)],
            max_val     = row[col.get("Max", 8)],
            default_val = row[col.get("Default", 9)],
        )
    wb.close()
    return configs
