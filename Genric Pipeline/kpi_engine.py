"""
Generic KPI engine for all SEUs (reverse-engineered from the EO feature file).

Idea (works for ANY equipment / KPI):
  - PI_Database_EO.xlsx   : PI Sensor / generic Tag Name -> Value   (the leaves)
  - feature_file `inferred`: Tag Name -> formula_expression          (the calc graph)

Any KPI is just an inferred tag. To compute it we recursively resolve its
`[tag]` references: if a ref has an inferred formula we evaluate it, otherwise
it's a leaf and we read its value from the PI database. The dependency chains
bottom out entirely in the PI database (verified), so this evaluates fully.

Same engine for every SEU — only the list of target KPI tags differs.
"""
from __future__ import annotations

import re
import math
import openpyxl
from pathlib import Path

REF = re.compile(r"\[([^\]]+)\]")

_SAFE = {
    "_if": lambda c, a, b: a if c else b,
    "min": min, "max": max, "abs": abs,
    "exp": math.exp, "log": math.log, "sqrt": math.sqrt,
    "pow": pow,
    "missing": lambda x: x is None,        # inferred-sheet null-guard
}


def _to_float(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


class KPIEngine:
    def __init__(self, pi_db_path: str, feature_path: str):
        self.values = self._load_pi_db(pi_db_path)      # leaf tag/sensor -> float
        self.formulas = self._load_inferred(feature_path)  # tag -> formula str
        self._cache: dict[str, float | None] = {}
        self.missing: set[str] = set()

    # ── loaders ──────────────────────────────────────────────────────────────
    @staticmethod
    def _load_pi_db(path):
        ws = openpyxl.load_workbook(path, data_only=True, read_only=True)["master_pi"]
        vals = {}
        for r in list(ws.iter_rows(values_only=True))[1:]:
            name, sensor, value = r[0], r[1], r[2]
            fv = _to_float(value)
            if name:   vals[str(name).strip()] = fv
            if sensor: vals[str(sensor).strip()] = fv
        return vals

    @staticmethod
    def _load_inferred(path):
        ws = openpyxl.load_workbook(path, data_only=True, read_only=True)["inferred"]
        return {str(r[0]).strip(): str(r[1])
                for r in list(ws.iter_rows(values_only=True))[1:]
                if r[0] and r[1] is not None}

    # ── evaluation ───────────────────────────────────────────────────────────
    def eval(self, tag: str, _stack: frozenset = frozenset()) -> float | None:
        tag = tag.strip()
        if tag in self._cache:
            return self._cache[tag]
        if tag in _stack:                      # cycle guard
            return None
        if tag not in self.formulas:           # leaf -> PI database
            v = self.values.get(tag)
            if v is None and tag not in self.values:
                self.missing.add(tag)
            self._cache[tag] = v
            return v
        # derived -> resolve refs then evaluate
        formula = self.formulas[tag]
        py = formula
        for ref in set(REF.findall(formula)):
            val = self.eval(ref, _stack | {tag})
            py = py.replace(f"[{ref}]", "(" + repr(val) + ")")
        py = py.replace("||", " or ").replace("&&", " and ")
        py = py.replace("^", "**")            # inferred uses ^ for exponentiation
        py = re.sub(r"\bif\(", "_if(", py)
        try:
            result = eval(py, {"__builtins__": {}}, _SAFE)  # noqa: S307 (controlled input)
        except Exception:
            result = None
        result = _to_float(result)
        self._cache[tag] = result
        return result

    # ── leaf sets + KPI matching (generic) ───────────────────────────────────
    def leaves(self, tag: str, _stack: frozenset = frozenset()) -> set[str]:
        """Transitive set of leaf (PI-DB) tags a derived tag depends on."""
        tag = tag.strip()
        if not hasattr(self, "_leafcache"):
            self._leafcache = {}
        if tag in self._leafcache:
            return self._leafcache[tag]
        f = self.formulas.get(tag)
        if f is None:
            return {tag}
        if tag in _stack:
            return set()
        res: set[str] = set()
        for ref in set(REF.findall(f)):
            res |= self.leaves(ref, _stack | {tag})
        self._leafcache[tag] = res
        return res

    def match_kpi(self, input_names: set[str], topn: int = 3):
        """Find inferred tags whose dependency-leaves best cover the KPI inputs."""
        scored = []
        for tag in self.formulas:
            inter = self.leaves(tag) & input_names
            if not inter:
                continue
            coverage = len(inter) / max(len(input_names), 1)
            scored.append((round(coverage, 3), len(inter), tag))
        scored.sort(reverse=True)
        return scored[:topn]

    # ── trace (show the calculation) ─────────────────────────────────────────
    def trace(self, tag: str, depth: int = 0, max_depth: int = 4, _stack=frozenset()):
        tag = tag.strip()
        ind = "  " * depth
        if tag not in self.formulas:
            print(f"{ind}{tag} = {self.values.get(tag)}   [PI-DB leaf]")
            return
        if tag in _stack or depth >= max_depth:
            print(f"{ind}{tag} = {self.eval(tag)}   (...)")
            return
        print(f"{ind}{tag} = {self.eval(tag)}")
        print(f"{ind}   := {self.formulas[tag]}")
        for ref in dict.fromkeys(REF.findall(self.formulas[tag])):
            self.trace(ref, depth + 1, max_depth, _stack | {tag})


if __name__ == "__main__":
    ROOT = Path(r"C:\Users\tnigam\Desktop\Python EO\Genric Pipeline\Data\source")
    eng = KPIEngine(ROOT / "PI_Database_EO.xlsx", ROOT / "feature_file_eo_v9_unified.xlsx")

    print("=" * 72)
    print("FURNACE a (furn1) — Thermal/Net Efficiency KPI, full calculation chain")
    print("=" * 72)
    eng.trace("furn1_net_efficiency", max_depth=5)
    print(f"\n>>> furn1_net_efficiency = {eng.eval('furn1_net_efficiency')}")
    if eng.missing:
        print(f"\n[missing leaves]: {sorted(eng.missing)}")
