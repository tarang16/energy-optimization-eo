"""
Build a single master PI sheet: Tag Name | PI Sensor | Value.

Merges two sources, keyed by the common PI Sensor ID:
  1) feature_file_eo_v9_unified.xlsx
        - `tag` sheet            : tag_name <-> pi_name (PI sensor)
        - `master_pi_data` sheet : tag_name -> real value (one snapshot row)
  2) Boiler_PEEO_Tags_with_results.xlsx
        - `pi_tags` sheet        : short name <-> pi_tags <-> value

Value precedence: real value from master_pi_data, else dictionary value.
Tag Name precedence: feature_file tag_name, else dictionary short name.
"""
from __future__ import annotations

from pathlib import Path
import openpyxl
from openpyxl import Workbook

FEATURE = Path(r"C:\Users\tnigam\Desktop\Python EO\Genric Pipeline\Data\source\feature_file_eo_v9_unified.xlsx")
DICT    = Path(r"C:\Users\tnigam\Desktop\Python EO\Boiler_PEEO_Tags_with_results.xlsx")
OUT     = Path(r"C:\Users\tnigam\Desktop\Python EO\Genric Pipeline\Data\source\master_pi_sheet.xlsx")


def load_feature():
    wb = openpyxl.load_workbook(FEATURE, data_only=True, read_only=True)
    tg = list(wb["tag"].iter_rows(values_only=True))           # tag_name, pi_name, tag_type, data_type
    mp = list(wb["master_pi_data"].iter_rows(values_only=True))
    hdr, val = mp[0], mp[1]
    name2val = {str(hdr[i]).strip(): val[i] for i in range(len(hdr))}
    pairs = []
    for r in tg[1:]:
        tag_name, pi_name = r[0], r[1]
        if not pi_name or not tag_name:
            continue
        pairs.append((str(tag_name).strip(), str(pi_name).strip(),
                      name2val.get(str(tag_name).strip())))
    return pairs


def load_dict():
    wb = openpyxl.load_workbook(DICT, data_only=True, read_only=True)
    s = list(wb["pi_tags"].iter_rows(values_only=True))         # short name[1], value[2], pi_tags[3]
    out = []
    for r in s[1:]:
        if r[3]:
            out.append((str(r[1]).strip() if r[1] else "", str(r[3]).strip(), r[2]))
    return out


def main():
    master: dict[str, dict] = {}   # pi_sensor -> {tag_name, value}

    # 1) feature_file first (authoritative naming + real values)
    for tag_name, pi, value in load_feature():
        m = master.setdefault(pi, {"tag_name": tag_name, "value": None})
        if value is not None:
            m["value"] = value

    # 2) dictionary: add missing sensors, and fill values where feature_file had none
    for short_name, pi, value in load_dict():
        m = master.get(pi)
        if m is None:
            master[pi] = {"tag_name": short_name, "value": value}
        elif m["value"] is None and value is not None:
            m["value"] = value

    rows = sorted(((v["tag_name"], pi, v["value"]) for pi, v in master.items()),
                  key=lambda x: x[0].lower())

    wb = Workbook(); ws = wb.active; ws.title = "master_pi"
    ws.append(["Tag Name", "PI Sensor", "Value"])
    for tn, pi, val in rows:
        ws.append([tn, pi, val])
    wb.save(OUT)

    total = len(rows)
    with_val = sum(1 for _, _, v in rows if v is not None)
    print(f"Wrote {OUT}")
    print(f"  rows (unique PI sensors): {total}")
    print(f"  with a value: {with_val}  |  empty: {total - with_val}")


if __name__ == "__main__":
    main()
