"""
GENERIC SEU KPI pipeline — instance/KPI names are irrelevant.

For each KPI in the SEU sheet it reads the unit's INPUT ATTRIBUTES by their
standard names, resolves them to PI-Database values, and applies the GENERIC
formula for that equipment + metric. ALL formulas live in kpi_formulas.py.

  Turbine / Exchanger -> pure input-driven functions (kpi_formulas)
  Furnace / Boiler    -> inferred [Tag Name] formulas (kpi_formulas.TAG_KPIS)

Output: Data/source/SEU_Sheet_KPIs.xlsx  (scoped to the unique SEU-sheet KPIs)
"""
from __future__ import annotations
import re, openpyxl
from string import ascii_lowercase
from collections import defaultdict
from pathlib import Path
from openpyxl import Workbook
from kpi_engine import KPIEngine
import kpi_formulas as F

ROOT = Path(r"C:\Users\tnigam\Desktop\Python EO\Genric Pipeline\Data\source")
EXPORT = Path(r"C:\Users\tnigam\Downloads\system_kpis_all_2026-06-04 (2).xlsx")
OUT = ROOT / "SEU_Sheet_KPIs.xlsx"


def _norm(s):
    return re.sub(r"[^a-z0-9]+", "", str(s).lower().replace("₂", "2"))


def main():
    eng = KPIEngine(ROOT / "PI_Database_EO.xlsx", ROOT / "feature_file_eo_v9_unified.xlsx")
    eta = eng.eval("HP_Turbine_Efficiency") or 0.8
    pdb = openpyxl.load_workbook(ROOT / "PI_Database_EO.xlsx", data_only=True, read_only=True)["master_pi"]
    pi2val = {str(r[1]).strip(): r[2] for r in list(pdb.iter_rows(values_only=True))[1:] if r[1]}

    def num(v):
        try: return float(v)
        except (TypeError, ValueError): return None

    ws = openpyxl.load_workbook(EXPORT, data_only=True, read_only=True)["SEU KPIs"]
    rows = list(ws.iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(rows) if r and r[0] == "Plant")
    h = {v: i for i, v in enumerate(rows[hi]) if v}; ci = lambda k: h[k]

    inst_inputs: dict = {}
    uniq, seen = [], set()
    for r in rows[hi + 1:]:
        if not r or not r[ci("KPI Name")] or not r[ci("Element")]:
            continue
        key = (str(r[ci("Plant")]).strip(), str(r[ci("Element")]).strip(), str(r[ci("Instance")]).strip())
        if r[ci("Attribute Name")] and r[ci("PI Tag")]:
            inst_inputs.setdefault(key, {}).setdefault(
                str(r[ci("Attribute Name")]).strip(), num(pi2val.get(str(r[ci("PI Tag")]).split()[0])))
        ukey = key + (str(r[ci("Network")]).strip(), str(r[ci("KPI Name")]).strip())
        if ukey not in seen:
            seen.add(ukey); uniq.append(ukey)

    # tag-based metric lookup (Furnace/Boiler) from kpi_formulas, normalised labels
    tag_lookup = {(e, _norm(lbl)): spec
                  for e, (_letters, kpis) in F.TAG_KPIS.items() for lbl, spec in kpis.items()}

    def resolve(plant, elem, inst, kpi):
        metric = kpi.split(" - ")[-1]
        inp = inst_inputs.get((plant, elem, inst), {})
        if "Turbine" in elem:
            return F.turbine_kpi(metric, inp, eta)
        if elem == "Steam Exchanger":
            return F.exchanger_kpi(metric, inp)
        if elem in F.TAG_KPIS:
            t = inst.strip().split()[-1]
            idx = (ascii_lowercase.index(t) + 1) if len(t) == 1 and t in ascii_lowercase else None
            spec = tag_lookup.get((elem, _norm(metric)))
            if idx and spec is not None:
                return F.value_for(eng, spec, idx)[0]
        return None

    out = Workbook(); o = out.active; o.title = "seu_sheet_kpis"
    o.append(["Plant", "Network", "Element", "Instance", "KPI Name", "Value", "Status"])
    agg = defaultdict(lambda: [0, 0]); n_val = 0
    for plant, elem, inst, net, kpi in uniq:
        v = resolve(plant, elem, inst, kpi)
        o.append([plant, net, elem, inst, kpi, round(v, 3) if isinstance(v, float) else v,
                  "CALCULATED" if v is not None else "no formula/input"])
        agg[elem][0] += 1; agg[elem][1] += (v is not None); n_val += (v is not None)
    out.save(OUT)
    print(f"Wrote {OUT}\n  unique SEU-sheet KPIs: {len(uniq)} | calculated: {n_val}")
    for e, (t, c) in sorted(agg.items()):
        print(f"    {e:<32}{c}/{t}")


if __name__ == "__main__":
    main()
