"""
Build one KPI "package" per equipment category from the SEU registry.

Output: Data/source/SEU_KPI_Packages.xlsx
    - one sheet per seu_category (Furnace, Boiler, Compressor, ...) listing each
      unit and its computed KPIs (Actual/Baseline/Target Duty, ENPI, Gain)
    - a Summary sheet with coverage per category
"""
from __future__ import annotations
from pathlib import Path
from openpyxl import Workbook
from seu_engine import SEUEngine, EXPR_COLS
import seu_mappings as M

OUT = Path(r"C:\Users\tnigam\Desktop\Python EO\Genric Pipeline\Data\source\SEU_KPI_Packages.xlsx")
KPI_LABELS = list(EXPR_COLS.values())


def _fmt(v):
    return round(v, 3) if isinstance(v, float) else v


def main():
    se = SEUEngine()
    M.apply(se)                       # register composed turbine KPI formulas (mapping file)
    pkgs = se.compute()

    wb = Workbook()
    summ = wb.active; summ.title = "Summary"
    summ.append(["Equipment (seu_category)", "Units", "Actual-Duty computed", "Coverage %"])

    grand_n = grand_d = 0
    for cat in sorted(pkgs):
        units = pkgs[cat]
        ws = wb.create_sheet(cat[:31])
        ws.append(["SEU", "Display Name", "Energy"] + KPI_LABELS)
        d = 0
        for u in units:
            ws.append([u["seu"], u["display"], u["energy"]] + [_fmt(u[k]) for k in KPI_LABELS])
            if u["Actual Duty (GJ/h)"] is not None:
                d += 1
        n = len(units); grand_n += n; grand_d += d
        summ.append([cat, n, d, round(100 * d / n) if n else 0])
    summ.append(["TOTAL", grand_n, grand_d, round(100 * grand_d / grand_n)])

    # composed turbine KPIs (from seu_mappings) — power / steam-rate / isentropic eff
    tk = wb.create_sheet("Turbines (composed)")
    tk.append(["KPI", "Value"])
    tcomp = 0
    for label, tag in M.turbine_kpi_outputs().items():
        v = se.eng.eval(tag)
        tk.append([label, _fmt(v)])
        if v is not None:
            tcomp += 1
    wb.save(OUT)
    main.tcomp = tcomp

    # console preview
    print(f"Wrote {OUT}\n")
    for cat in ["Furnace", "Boiler", "Compressor", "Exchanger"]:
        print(f"=== {cat} package ===")
        for u in pkgs[cat][:4]:
            print(f"  {u['seu']:<10} {u['display'][:28]:<30} "
                  f"ActualDuty={_fmt(u['Actual Duty (GJ/h)'])}  ENPI={_fmt(u['ENPI'])}")
        print()
    print(f"\nSEU registry duties computed : {grand_d}/{grand_n}")
    print(f"Composed turbine KPIs computed: {getattr(main, 'tcomp', 0)}")
    print(f"ONE package -> {OUT.name}")


if __name__ == "__main__":
    main()
