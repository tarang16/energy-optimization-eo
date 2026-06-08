import pandas as pd, re, math, sys
from pathlib import Path
sys.stdout.reconfigure(encoding='utf-8')

FF   = Path(r"C:\Users\tnigam\Desktop\Python EO\Genric Pipeline\Data\source\feature_file_eo_v9_unified.xlsx")
HIDX = Path(r"C:\Users\tnigam\Desktop\Python EO\Genric Pipeline\Data\source\_hierarchy_index.csv")
OUT  = Path(r"C:\Users\tnigam\Desktop\Python EO\Genric Pipeline\eo_pipeline\docs\boiler_inferred_calc_results.xlsx")

BOILER_CATEGORIES = {
    "Boiler -- Combustion & Fuel","Boiler -- Performance & Energy",
    "Boiler -- Flue Gas & Stack","Boiler -- Air System",
    "Boiler -- BFW System","Boiler -- Steam Generation",
    "Boiler -- Blowdown","Boiler -- Emissions",
}
REF_RE = re.compile(r"\[([^\[\]]+)\]")

# ── Load ──────────────────────────────────────────────────────────────────────
tag_df  = pd.read_excel(FF, sheet_name='tag')
inf_df  = pd.read_excel(FF, sheet_name='inferred')
mpd     = pd.read_excel(FF, sheet_name='master_pi_data')
hier_df = pd.read_csv(HIDX)

mpd_row  = mpd.iloc[0].to_dict()
mpd_date = str(mpd_row.get('date',''))

# ── Hierarchy-covered L2 tags ─────────────────────────────────────────────────
hier_sensors = set(hier_df[hier_df['has_pi_sensor']==True]['pi_sensors'].dropna().astype(str))
l2_to_pi     = dict(zip(tag_df['tag_name'].astype(str), tag_df['pi_name'].astype(str)))
hierarchy_l2  = {l2 for l2,pi in l2_to_pi.items() if str(pi) in hier_sensors}

# seed values: ONLY hierarchy-sourced PI tags
base_vals: dict = {}
for tag in hierarchy_l2:
    v = mpd_row.get(tag)
    if v is not None:
        try:
            fv = float(v)
            if not math.isnan(fv):
                base_vals[tag] = fv
        except:
            pass

# ── Inferred maps ─────────────────────────────────────────────────────────────
inf_records = inf_df.dropna(subset=['tag_name']).drop_duplicates('tag_name')
inf_map  = inf_records.set_index('tag_name')['formula_expression'].to_dict()
cat_map  = inf_records.set_index('tag_name')['Category'].to_dict()
boiler_tags = sorted(t for t,c in cat_map.items() if c in BOILER_CATEGORIES)

print(f"Hierarchy sensors: {len(hier_sensors)}  |  L2 tags with values: {len(base_vals)}")
print(f"Boiler inferred targets: {len(boiler_tags)}")

# ── Case-insensitive tag lookup ───────────────────────────────────────────────
all_known = set(base_vals.keys()) | set(inf_map.keys())
upper_to_actual: dict = {t.upper(): t for t in all_known}

def resolve_tag(name):
    if name in base_vals or name in inf_map:
        return name
    return upper_to_actual.get(name.upper())

# ── Eval helpers ──────────────────────────────────────────────────────────────
def _missing(v):
    return v is None or (isinstance(v, float) and math.isnan(v))

EVAL_NS = {
    'if_':    lambda c,a,b: a if c else b,
    'missing': _missing,
    'abs': abs, 'round': round,
    'max': max, 'min': min,
    'sqrt': math.sqrt, 'log': math.log, 'exp': math.exp,
    'ceil': math.ceil, 'floor': math.floor,
    'nan': float('nan'), 'inf': float('inf'),
    '__builtins__': {},
}

def to_py(formula, vals):
    s = str(formula).replace('^', '**')
    s = re.sub(r'\bif\s*\(', 'if_(', s)
    missing_refs = []

    def replace_ref(m):
        raw   = m.group(1)
        canon = resolve_tag(raw)
        if canon is not None and canon in vals:
            return f"({vals[canon]!r})"
        missing_refs.append(raw)
        return "float('nan')"

    s = REF_RE.sub(replace_ref, s)
    return s, missing_refs

# ── Recursive evaluator ───────────────────────────────────────────────────────
memo = {}

def evaluate(tag, live_vals, stack):
    if tag in memo:
        return memo[tag]
    if tag in stack:
        r = (None, 'cycle_detected', [tag])
        memo[tag] = r
        return r
    stack.add(tag)

    canon = resolve_tag(tag) or tag
    if canon in live_vals:
        r = (live_vals[canon], 'hierarchy_pi_value', [])
        memo[tag] = r
        stack.discard(tag)
        return r

    formula = inf_map.get(tag)
    if formula is None or (isinstance(formula, float) and math.isnan(formula)):
        r = (None, 'no_formula_no_pi', [tag])
        memo[tag] = r
        stack.discard(tag)
        return r

    fstr = str(formula)
    refs = set(REF_RE.findall(fstr))
    for ref in refs:
        if ref == tag:
            continue
        canon_ref = resolve_tag(ref)
        if canon_ref and canon_ref not in live_vals:
            rv, _, _ = evaluate(canon_ref, live_vals, stack)
            if rv is not None:
                live_vals[canon_ref] = rv
                upper_to_actual[canon_ref.upper()] = canon_ref

    py_expr, missing = to_py(fstr, live_vals)

    if missing:
        r = (None, 'missing_inputs', sorted(set(missing)))
        memo[tag] = r
        stack.discard(tag)
        return r

    try:
        v = eval(py_expr, dict(EVAL_NS))
        if isinstance(v, bool):
            v = int(v)
        v = float(v)
        if math.isnan(v):
            r = (None, 'result_is_nan', [])
        else:
            r = (v, 'computed', [])
    except ZeroDivisionError:
        r = (None, 'div_by_zero', [])
    except SyntaxError as e:
        r = (None, f'syntax_error: {str(e)[:80]}', [])
    except Exception as e:
        r = (None, f'eval_error: {str(e)[:80]}', [])

    memo[tag] = r
    stack.discard(tag)
    return r

# ── Run evaluations ───────────────────────────────────────────────────────────
live_vals = dict(base_vals)
rows = []
stack = set()

for tag in boiler_tags:
    val, status, missing = evaluate(tag, live_vals, stack)
    if val is not None:
        live_vals[tag] = val

    formula = inf_map.get(tag, '')
    if isinstance(formula, float) and math.isnan(formula):
        formula = ''

    reason = ''
    if status == 'missing_inputs':
        listed = missing[:6]
        reason = 'Inputs not in hierarchy: ' + '; '.join(listed)
        if len(missing) > 6:
            reason += f' ... (+{len(missing)-6} more)'
    elif status not in ('computed', 'hierarchy_pi_value'):
        reason = status

    rows.append({
        'tag_name'         : tag,
        'category'         : cat_map.get(tag, ''),
        'formula'          : str(formula)[:300],
        'computed_status'  : 'COMPUTED' if status in ('computed','hierarchy_pi_value') else 'NOT COMPUTED',
        'value'            : round(val, 6) if val is not None else '',
        'reason_if_not'    : reason,
        'missing_inputs'   : '; '.join(missing) if status == 'missing_inputs' else '',
    })

df = pd.DataFrame(rows)

# ── Write Excel ───────────────────────────────────────────────────────────────
computed     = df[df['computed_status']=='COMPUTED'][['tag_name','category','formula','value']].copy()
not_computed = df[df['computed_status']=='NOT COMPUTED'][['tag_name','category','formula','reason_if_not','missing_inputs']].copy()

with pd.ExcelWriter(OUT, engine='openpyxl') as writer:
    computed.to_excel(writer,     sheet_name='Computed',     index=False)
    not_computed.to_excel(writer, sheet_name='Not_Computed', index=False)
    df.to_excel(writer,           sheet_name='All_Results',  index=False)

    from openpyxl.styles import PatternFill, Font
    green = PatternFill("solid", fgColor="C6EFCE")
    red   = PatternFill("solid", fgColor="FFC7CE")

    # Colour All_Results sheet
    ws = writer.sheets['All_Results']
    status_col = None
    for cell in ws[1]:
        if cell.value == 'computed_status':
            status_col = cell.column
            break
    if status_col:
        for row in ws.iter_rows(min_row=2):
            cell = row[status_col-1]
            fill = green if cell.value == 'COMPUTED' else red
            for c in row:
                c.fill = fill

    # Auto-width all sheets
    for sh_name, ws in writer.sheets.items():
        for col in ws.columns:
            max_len = max((len(str(c.value or '')) for c in col), default=10)
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 55)

print(f"Wrote: {OUT}")
print(f"\n{'='*56}")
print(f" BOILER CALC RESULTS  (data date: {mpd_date})")
print(f" Input: {len(base_vals)} hierarchy-sourced PI values")
print(f"{'='*56}")
sc = df['computed_status'].value_counts()
for s,n in sc.items():
    print(f"  {s:<14s}  {n:4d}  ({100*n/len(df):.1f}%)")

print(f"\n── Sample COMPUTED tags ──")
for _,r in computed.head(12).iterrows():
    print(f"  {r.tag_name:<50s}  {r.value}")

print(f"\n── NOT COMPUTED — reason summary ──")
rsummary = not_computed['reason_if_not'].str[:40].value_counts().head(8)
for r,n in rsummary.items():
    print(f"  {n:4d}  {r}")

print(f"\n── Top missing inputs blocking most tags ──")
miss = (not_computed[not_computed['missing_inputs'] != '']
        ['missing_inputs'].str.split('; ').explode()
        .str.strip().loc[lambda s: s != ''].value_counts())
for tag,n in miss.head(15).items():
    print(f"  {n:3d}  {tag}")
