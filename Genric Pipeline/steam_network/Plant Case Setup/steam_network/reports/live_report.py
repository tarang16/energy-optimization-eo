"""Generate live Word SOP + Excel Report from current network state.

Called automatically after every successful solve.  Both documents are written
to the configured output directory and served via GET endpoints.

The reports contain:
  * Actual component specs (from the live graph engine)
  * Actual solve results (flows, pressures, enthalpies, power)
  * Per-header balance totals
  * Solver convergence log
  * Thermodynamic state points at each component
"""
from __future__ import annotations

import datetime
import traceback
from pathlib import Path
from typing import Any, Optional

from ..core.graph_engine import GraphEngine
from ..core.logger import get_logger
from ..models.enums import ComponentType, BalanceStatus
from ..models.schemas import SolveResult

log = get_logger("reports")

# ---------- default output dir ----------
REPORT_DIR = Path(__file__).resolve().parent.parent.parent / "reports_output"

# ---------- helpers ----------

def _safe(val, fmt=".1f"):
    if val is None:
        return "--"
    try:
        return f"{val:{fmt}}"
    except Exception:
        return str(val)


def _ensure_dir(d: Path):
    d.mkdir(parents=True, exist_ok=True)


# ============================================================
#  PUBLIC API
# ============================================================

def generate_reports(
    engine: GraphEngine,
    solve_result: Optional[SolveResult],
    output_dir: Optional[Path] = None,
) -> dict[str, Path]:
    """Generate both Word and Excel reports.  Returns {"docx": path, "xlsx": path}."""
    out = output_dir or REPORT_DIR
    _ensure_dir(out)
    paths: dict[str, Path] = {}
    try:
        paths["xlsx"] = _generate_excel(engine, solve_result, out)
    except Exception:
        log.error("Excel report generation failed:\n%s", traceback.format_exc())
    try:
        paths["docx"] = _generate_word(engine, solve_result, out)
    except Exception:
        log.error("Word report generation failed:\n%s", traceback.format_exc())
    return paths


# ============================================================
#  EXCEL REPORT
# ============================================================

def _generate_excel(engine: GraphEngine, result: Optional[SolveResult], out: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    now = datetime.datetime.now()

    # -- styles --
    HDR_FONT = Font(name="Calibri", bold=True, size=11, color="FFFFFF")
    HDR_FILL = PatternFill(start_color="1A3C6E", end_color="1A3C6E", fill_type="solid")
    SUB_FONT = Font(name="Calibri", bold=True, size=11, color="1A3C6E")
    SUB_FILL = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
    NORM = Font(name="Calibri", size=11)
    BORDER = Border(left=Side("thin"), right=Side("thin"),
                    top=Side("thin"), bottom=Side("thin"))
    OK_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    WARN_FILL = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
    BAD_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

    def hdr_row(ws, row, vals):
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=row, column=c, value=v)
            cell.font = HDR_FONT; cell.fill = HDR_FILL
            cell.border = BORDER
            cell.alignment = Alignment(horizontal="center", wrap_text=True)

    def sub_row(ws, row, vals):
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=row, column=c, value=v)
            cell.font = SUB_FONT; cell.fill = SUB_FILL
            cell.border = BORDER

    def data_row(ws, row, vals, fmt=None, fill=None):
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=row, column=c, value=v)
            cell.font = NORM; cell.border = BORDER
            cell.alignment = Alignment(horizontal="center" if isinstance(v, (int, float)) else "left")
            if fmt and isinstance(v, (int, float)):
                cell.number_format = fmt
            if fill:
                cell.fill = fill

    def auto_w(ws):
        for col in ws.columns:
            mx = 0
            letter = get_column_letter(col[0].column)
            for cell in col:
                try:
                    if cell.value:
                        mx = max(mx, len(str(cell.value)))
                except Exception:
                    pass
            ws.column_dimensions[letter].width = min(mx + 4, 45)

    # Gather live data
    topo = engine.to_dict()
    nodes = topo["nodes"]
    edges = topo["edges"]
    state_map = {}
    if result:
        state_map = {s.id: s for s in result.component_states}

    # ===== Sheet 1: Solve Summary ===================================
    ws = wb.active
    ws.title = "Solve Summary"
    ws.merge_cells("A1:F1")
    ws["A1"] = "STEAM NETWORK - LIVE SOLVE REPORT"
    ws["A1"].font = Font(name="Calibri", bold=True, size=14, color="1A3C6E")
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.merge_cells("A2:F2")
    ws["A2"] = f"Generated: {now.strftime('%Y-%m-%d %H:%M:%S')}"
    ws["A2"].font = Font(name="Calibri", size=10, italic=True, color="666666")
    ws["A2"].alignment = Alignment(horizontal="center")

    r = 4
    hdr_row(ws, r, ["Parameter", "Value", "Unit", "Notes"])
    summary = [
        ("Total Components", len(nodes), "", ""),
        ("Total Connections", len(edges), "", ""),
        ("Solver Status", result.status.value if result else "unsolved", "",
         "ok = converged, unbalanced = stagnated"),
        ("Iterations", result.iterations if result else 0, "", ""),
        ("Mass Residual", result.residual_mass_tph if result else None, "t/h",
         "< 0.001 = converged"),
        ("Energy Residual", result.residual_energy_kw if result else None, "kW",
         "< 0.1 = converged"),
    ]
    for i, (p, v, u, n) in enumerate(summary):
        fill = None
        if p == "Solver Status":
            fill = OK_FILL if v == "ok" else (WARN_FILL if v == "unbalanced" else BAD_FILL)
        data_row(ws, r+1+i, [p, v, u, n], fmt="#,##0.0000", fill=fill)

    # Solver messages
    r += len(summary) + 2
    if result and result.messages:
        sub_row(ws, r, ["Solver Messages", "", "", ""])
        for i, msg in enumerate(result.messages):
            data_row(ws, r+1+i, [msg, "", "", ""])
        r += len(result.messages) + 1

    auto_w(ws)

    # ===== Sheet 2: Component Inventory =============================
    ws2 = wb.create_sheet("Component Inventory")
    r = 1
    hdr_row(ws2, r, ["ID", "Name", "Type", "Header Level", "Key Parameter", "Value", "Unit"])

    r = 2
    for n in sorted(nodes, key=lambda x: x["type"]):
        lvl = ""
        key_p = ""
        val = ""
        unit = ""
        spec = (n.get("metadata") or {}).get("spec", {})
        ports = n.get("ports", [])
        for p in ports:
            if p.get("header_level"):
                lvl = p["header_level"]
                break
        t = n["type"]
        if t == "source":
            key_p = "Capacity"; val = spec.get("capacity_tph", ""); unit = "t/h"
        elif t == "consumer":
            key_p = "Demand"; val = spec.get("demand_tph", ""); unit = "t/h"
        elif t == "turbine":
            key_p = "Inlet Flow"; val = spec.get("inlet_flow_tph", ""); unit = "t/h"
        elif t in ("prds", "valve"):
            key_p = "Max Capacity"; val = spec.get("max_capacity_tph", spec.get("max_flow_tph", "")); unit = "t/h"
        elif t == "header":
            key_p = "Pressure"
            val = ports[0].get("nominal_pressure_bar", "") if ports else ""
            unit = "bar"
        elif t == "condenser":
            key_p = "Capacity"; val = spec.get("capacity_tph", ""); unit = "t/h"
        data_row(ws2, r, [n["id"], n["name"], t, lvl, key_p, val, unit], fmt="#,##0.0")
        r += 1
    auto_w(ws2)

    # ===== Sheet 3: Component States (post-solve) ===================
    ws3 = wb.create_sheet("Solved States")
    r = 1
    hdr_row(ws3, r, [
        "Component", "Type",
        "Inlet P (bar)", "Inlet T (C)", "Inlet h (kJ/kg)", "Inlet m (t/h)",
        "Outlet P (bar)", "Outlet T (C)", "Outlet h (kJ/kg)", "Outlet m (t/h)",
        "Power (kW)", "Duty (kW)",
    ])
    r = 2
    for s in (result.component_states if result else []):
        si = s.inlet
        so = s.outlet
        data_row(ws3, r, [
            s.name, s.type.value,
            si.pressure_bar if si else None,
            si.temperature_c if si else None,
            si.enthalpy_kj_kg if si else None,
            si.mass_flow_tph if si else None,
            so.pressure_bar if so else None,
            so.temperature_c if so else None,
            so.enthalpy_kj_kg if so else None,
            so.mass_flow_tph if so else None,
            s.power_kw,
            s.duty_kw,
        ], fmt="#,##0.00")
        # Extraction states for turbines
        if s.extraction_states:
            for port_name, ex in s.extraction_states.items():
                data_row(ws3, r+1, [
                    f"  {s.name}.{port_name}", "extraction",
                    "", "", "", "",
                    ex.pressure_bar, ex.temperature_c,
                    ex.enthalpy_kj_kg, ex.mass_flow_tph,
                    "", "",
                ], fmt="#,##0.00")
                r += 1
        r += 1
    auto_w(ws3)

    # ===== Sheet 4: Header Balance ==================================
    ws4 = wb.create_sheet("Header Balance")
    r = 1
    hdr_row(ws4, r, [
        "Header", "Pressure (bar)", "Temperature (C)",
        "Generation (t/h)", "Demand (t/h)",
        "Imbalance (t/h)", "Status",
    ])

    # Compute per-header gen/dem from ACTUAL EDGES (connections), not port labels.
    # This is robust when a component's spec inlet_level differs from the header
    # it's actually connected to (e.g., turbine spec says VHP but wired to HP).
    header_totals: dict[str, dict[str, float]] = {}
    header_info: dict[str, dict] = {}
    header_id_to_lvl: dict[str, str] = {}
    for n in nodes:
        if n["type"] == "header":
            ports = n.get("ports", [])
            if ports:
                lvl = ports[0].get("header_level", "")
                press = ports[0].get("nominal_pressure_bar", 0)
                header_info[lvl] = {"pressure": press, "id": n["id"]}
                header_totals[lvl] = {"gen": 0.0, "dem": 0.0}
                header_id_to_lvl[n["id"]] = lvl

    if result:
        for e in edges:
            src_id = e["from"]
            tgt_id = e["to"]
            from_port = e.get("from_port", "")
            src_state = state_map.get(src_id)
            tgt_state = state_map.get(tgt_id)

            # Edge INTO a header = generation for that header
            tgt_lvl = header_id_to_lvl.get(tgt_id)
            if tgt_lvl and src_state:
                flow = 0.0
                if from_port == "out" or from_port == "exhaust":
                    flow = (src_state.outlet.mass_flow_tph or 0) if src_state.outlet else 0
                elif from_port.startswith("extraction_") and src_state.extraction_states:
                    ex = src_state.extraction_states.get(from_port)
                    flow = (ex.mass_flow_tph or 0) if ex else 0
                elif src_state.outlet:
                    flow = src_state.outlet.mass_flow_tph or 0
                if flow > 0:
                    header_totals[tgt_lvl]["gen"] += flow

            # Edge FROM a header = demand on that header
            src_lvl = header_id_to_lvl.get(src_id)
            if src_lvl and tgt_state:
                flow = (tgt_state.inlet.mass_flow_tph or 0) if tgt_state.inlet else 0
                if flow > 0:
                    header_totals[src_lvl]["dem"] += flow

    LEVEL_ORDER = ["VHP", "HP", "MP", "LP", "LLP", "LLP1"]
    r = 2
    for lvl in LEVEL_ORDER:
        if lvl not in header_totals:
            continue
        info = header_info.get(lvl, {})
        gen = header_totals[lvl]["gen"]
        dem = header_totals[lvl]["dem"]
        imb = gen - dem
        status = "Balanced" if abs(imb) < 0.1 else ("Surplus" if imb > 0 else "Deficit")
        fill = OK_FILL if abs(imb) < 0.1 else (WARN_FILL if abs(imb) < 5 else BAD_FILL)

        # get temperature from solved state of the header node
        temp = None
        hid = info.get("id")
        if hid and hid in state_map:
            hs = state_map[hid]
            if hs.outlet and hs.outlet.temperature_c:
                temp = hs.outlet.temperature_c
            elif hs.inlet and hs.inlet.temperature_c:
                temp = hs.inlet.temperature_c

        data_row(ws4, r, [
            lvl, info.get("pressure"), temp,
            round(gen, 2), round(dem, 2), round(imb, 2), status,
        ], fmt="#,##0.00", fill=fill)
        r += 1

    # Totals row
    total_gen = sum(ht["gen"] for ht in header_totals.values())
    total_dem = sum(ht["dem"] for ht in header_totals.values())
    r += 1
    sub_row(ws4, r, ["TOTAL", "", "", f"{total_gen:.1f}", f"{total_dem:.1f}",
                      f"{total_gen - total_dem:.2f}", ""])
    auto_w(ws4)

    # ===== Sheet 5: Connection Details ==============================
    ws5 = wb.create_sheet("Connections")
    r = 1
    hdr_row(ws5, r, ["From Component", "From Port", "To Component", "To Port",
                      "Pressure (bar)", "Nominal Flow (t/h)"])
    r = 2
    for e in edges:
        from_name = next((n["name"] for n in nodes if n["id"] == e["from"]), e["from"])
        to_name = next((n["name"] for n in nodes if n["id"] == e["to"]), e["to"])
        data_row(ws5, r, [
            from_name, e["from_port"], to_name, e["to_port"],
            e.get("pressure_bar"), e.get("nominal_flow_tph"),
        ], fmt="#,##0.0")
        r += 1
    auto_w(ws5)

    # ===== Sheet 6: Source Details ==================================
    ws6 = wb.create_sheet("Sources Detail")
    r = 1
    hdr_row(ws6, r, ["Name", "Header", "Capacity (t/h)", "Min Load (t/h)",
                      "Pressure (bar)", "Temp (C)", "Efficiency",
                      "Actual Flow (t/h)", "Solved Flow (t/h)"])
    r = 2
    for n in nodes:
        if n["type"] != "source":
            continue
        spec = (n.get("metadata") or {}).get("spec", {})
        s = state_map.get(n["id"])
        solved_flow = s.outlet.mass_flow_tph if (s and s.outlet) else None
        data_row(ws6, r, [
            n["name"],
            spec.get("header_level", ""),
            spec.get("capacity_tph"),
            spec.get("min_load_tph"),
            spec.get("pressure_bar"),
            spec.get("temperature_c"),
            spec.get("efficiency"),
            spec.get("actual_flow_tph"),
            solved_flow,
        ], fmt="#,##0.00")
        r += 1
    auto_w(ws6)

    # ===== Sheet 7: Turbine Details =================================
    ws7 = wb.create_sheet("Turbines Detail")
    r = 1
    hdr_row(ws7, r, ["Name", "Mode", "Inlet Level", "Inlet Flow (t/h)",
                      "Extraction Levels", "Extraction Flows (t/h)",
                      "Exhaust Level", "Isen Eff", "Shaft Power (kW)",
                      "Gross Elec (kW)", "Solved Inlet (t/h)"])
    r = 2
    for n in nodes:
        if n["type"] != "turbine":
            continue
        spec = (n.get("metadata") or {}).get("spec", {})
        s = state_map.get(n["id"])
        shaft = s.metadata.get("shaft_power_kw") if (s and s.metadata) else None
        gross = s.metadata.get("gross_electrical_kw") if (s and s.metadata) else None
        solved_in = s.inlet.mass_flow_tph if (s and s.inlet) else None
        ext_lvls = ", ".join(spec.get("extraction_levels", []))
        ext_flows = ", ".join(str(f) for f in spec.get("extraction_flows_tph", []))
        data_row(ws7, r, [
            n["name"], spec.get("mode"), spec.get("inlet_level"),
            spec.get("inlet_flow_tph"), ext_lvls, ext_flows,
            spec.get("exhaust_level", "condenser"),
            spec.get("isentropic_efficiency"),
            shaft, gross, solved_in,
        ], fmt="#,##0.00")
        r += 1
    auto_w(ws7)

    # ===== Sheet 8: PRDS Details ====================================
    ws8 = wb.create_sheet("PRDS Detail")
    r = 1
    hdr_row(ws8, r, ["Name", "From Level", "To Level", "Max Cap (t/h)",
                      "Spray Temp (C)", "Solved Inlet (t/h)", "Solved Outlet (t/h)",
                      "Spray Flow (t/h)"])
    r = 2
    for n in nodes:
        if n["type"] != "prds":
            continue
        spec = (n.get("metadata") or {}).get("spec", {})
        s = state_map.get(n["id"])
        m_in = s.inlet.mass_flow_tph if (s and s.inlet) else None
        m_out = s.outlet.mass_flow_tph if (s and s.outlet) else None
        spray = (m_out - m_in) if (m_in is not None and m_out is not None) else None
        data_row(ws8, r, [
            n["name"], spec.get("from_level"), spec.get("to_level"),
            spec.get("max_capacity_tph"),
            spec.get("desuperheat_water_temp_c"),
            m_in, m_out, spray,
        ], fmt="#,##0.00")
        r += 1
    auto_w(ws8)

    # ===== Sheet 9: Consumer Details ================================
    ws9 = wb.create_sheet("Consumers Detail")
    r = 1
    hdr_row(ws9, r, ["Name", "Header", "Design Demand (t/h)",
                      "Return Fraction", "Return Temp (C)",
                      "Actual (PI tag)", "Solved Flow (t/h)"])
    r = 2
    for n in nodes:
        if n["type"] != "consumer":
            continue
        spec = (n.get("metadata") or {}).get("spec", {})
        s = state_map.get(n["id"])
        solved = s.inlet.mass_flow_tph if (s and s.inlet) else None
        data_row(ws9, r, [
            n["name"], spec.get("header_level"),
            spec.get("demand_tph"), spec.get("return_fraction"),
            spec.get("return_temperature_c"),
            spec.get("actual_flow_tph"), solved,
        ], fmt="#,##0.00")
        r += 1
    auto_w(ws9)

    fpath = out / "Steam_Network_Report.xlsx"
    wb.save(str(fpath))
    log.info("Excel report written: %s", fpath)
    return fpath


# ============================================================
#  WORD REPORT
# ============================================================

def _generate_word(engine: GraphEngine, result: Optional[SolveResult], out: Path) -> Path:
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    for lv in range(1, 5):
        doc.styles[f"Heading {lv}"].font.color.rgb = RGBColor(0x1A, 0x3C, 0x6E)

    now = datetime.datetime.now()
    topo = engine.to_dict()
    nodes = topo["nodes"]
    edges = topo["edges"]
    state_map = {}
    if result:
        state_map = {s.id: s for s in result.component_states}

    def add_styled_table(doc, headers):
        tbl = doc.add_table(rows=1, cols=len(headers), style="Light Grid Accent 1")
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(headers):
            tbl.rows[0].cells[i].text = h
            for p in tbl.rows[0].cells[i].paragraphs:
                for run in p.runs:
                    run.bold = True
        return tbl

    def tbl_row(tbl, vals):
        row = tbl.add_row()
        for i, v in enumerate(vals):
            row.cells[i].text = str(v) if v is not None else "--"
        return row

    # ---- COVER ----
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("STEAM NETWORK SIMULATION REPORT")
    run.bold = True; run.font.size = Pt(22)
    run.font.color.rgb = RGBColor(0x1A, 0x3C, 0x6E)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Live System Report\nAuto-generated from Network Solver")
    run.font.size = Pt(14); run.font.color.rgb = RGBColor(0x4A, 0x4A, 0x4A)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(
        f"Report Date: {now.strftime('%B %d, %Y %H:%M')}\n"
        f"Components: {len(nodes)}  |  Connections: {len(edges)}\n"
        f"Solver Status: {result.status.value if result else 'unsolved'}"
    )
    run.font.size = Pt(11)
    doc.add_page_break()

    # ---- 1. EXECUTIVE SUMMARY ----
    doc.add_heading("1. Network Summary", level=1)
    tbl = add_styled_table(doc, ["Metric", "Value"])
    tbl_row(tbl, ["Total Components", len(nodes)])
    tbl_row(tbl, ["Total Connections", len(edges)])
    tbl_row(tbl, ["Solver Status", result.status.value if result else "unsolved"])
    tbl_row(tbl, ["Iterations", result.iterations if result else 0])
    tbl_row(tbl, ["Mass Residual (t/h)", _safe(result.residual_mass_tph, ".4f") if result else "--"])
    tbl_row(tbl, ["Energy Residual (kW)", _safe(result.residual_energy_kw, ".2f") if result else "--"])

    # Count by type
    type_counts = {}
    for n in nodes:
        type_counts[n["type"]] = type_counts.get(n["type"], 0) + 1
    doc.add_paragraph()
    tbl = add_styled_table(doc, ["Component Type", "Count"])
    for t in sorted(type_counts.keys()):
        tbl_row(tbl, [t, type_counts[t]])

    if result and result.messages:
        doc.add_heading("Solver Messages", level=2)
        for msg in result.messages:
            doc.add_paragraph(msg, style="List Bullet")

    doc.add_page_break()

    # ---- 2. HEADER BALANCE ----
    doc.add_heading("2. Per-Header Mass Balance", level=1)
    doc.add_paragraph(
        "This table shows the total steam generation feeding each header and the "
        "total demand drawing from it. A balanced header has Imbalance close to zero."
    )

    # Compute header totals from ACTUAL EDGES (same edge-based logic as Excel + UI)
    header_totals: dict[str, dict[str, float]] = {}
    header_info: dict[str, dict] = {}
    header_id_to_lvl: dict[str, str] = {}
    for n in nodes:
        if n["type"] == "header":
            ports = n.get("ports", [])
            if ports:
                lvl = ports[0].get("header_level", "")
                press = ports[0].get("nominal_pressure_bar", 0)
                header_info[lvl] = {"pressure": press, "id": n["id"]}
                header_totals[lvl] = {"gen": 0.0, "dem": 0.0}
                header_id_to_lvl[n["id"]] = lvl

    if result:
        for e in edges:
            src_id = e["from"]
            tgt_id = e["to"]
            from_port = e.get("from_port", "")
            src_state = state_map.get(src_id)
            tgt_state = state_map.get(tgt_id)

            tgt_lvl = header_id_to_lvl.get(tgt_id)
            if tgt_lvl and src_state:
                flow = 0.0
                if from_port in ("out", "exhaust"):
                    flow = (src_state.outlet.mass_flow_tph or 0) if src_state.outlet else 0
                elif from_port.startswith("extraction_") and src_state.extraction_states:
                    ex = src_state.extraction_states.get(from_port)
                    flow = (ex.mass_flow_tph or 0) if ex else 0
                elif src_state.outlet:
                    flow = src_state.outlet.mass_flow_tph or 0
                if flow > 0:
                    header_totals[tgt_lvl]["gen"] += flow

            src_lvl = header_id_to_lvl.get(src_id)
            if src_lvl and tgt_state:
                flow = (tgt_state.inlet.mass_flow_tph or 0) if tgt_state.inlet else 0
                if flow > 0:
                    header_totals[src_lvl]["dem"] += flow

    tbl = add_styled_table(doc, ["Header", "Pressure (bar)", "Gen (t/h)", "Dem (t/h)", "Imbalance (t/h)", "Status"])
    LEVEL_ORDER = ["VHP", "HP", "MP", "LP", "LLP", "LLP1"]
    for lvl in LEVEL_ORDER:
        if lvl not in header_totals:
            continue
        gen = header_totals[lvl]["gen"]
        dem = header_totals[lvl]["dem"]
        imb = gen - dem
        status = "Balanced" if abs(imb) < 0.1 else ("Surplus" if imb > 0 else "Deficit")
        tbl_row(tbl, [lvl, _safe(header_info.get(lvl, {}).get("pressure")),
                       _safe(gen), _safe(dem), _safe(imb, ".2f"), status])

    doc.add_page_break()

    # ---- 3. SOURCES ----
    doc.add_heading("3. Steam Sources (Boilers / HRSGs)", level=1)
    src_nodes = [n for n in nodes if n["type"] == "source"]
    if src_nodes:
        tbl = add_styled_table(doc, ["Name", "Header", "Capacity (t/h)", "Pressure (bar)",
                                      "Temp (C)", "Efficiency", "Solved Flow (t/h)"])
        for n in src_nodes:
            spec = (n.get("metadata") or {}).get("spec", {})
            s = state_map.get(n["id"])
            flow = _safe(s.outlet.mass_flow_tph if (s and s.outlet) else None)
            tbl_row(tbl, [
                n["name"], spec.get("header_level"), spec.get("capacity_tph"),
                spec.get("pressure_bar"), spec.get("temperature_c"),
                spec.get("efficiency"), flow,
            ])
    else:
        doc.add_paragraph("No steam sources in the current network.")

    # ---- 4. TURBINES ----
    doc.add_heading("4. Turbines", level=1)
    turb_nodes = [n for n in nodes if n["type"] == "turbine"]
    if turb_nodes:
        for n in turb_nodes:
            spec = (n.get("metadata") or {}).get("spec", {})
            s = state_map.get(n["id"])
            doc.add_heading(f"Turbine: {n['name']}", level=2)
            tbl = add_styled_table(doc, ["Parameter", "Value", "Unit"])
            tbl_row(tbl, ["Mode", spec.get("mode"), ""])
            tbl_row(tbl, ["Inlet Level", spec.get("inlet_level"), ""])
            tbl_row(tbl, ["Design Inlet Flow", spec.get("inlet_flow_tph"), "t/h"])
            tbl_row(tbl, ["Solved Inlet Flow", _safe(s.inlet.mass_flow_tph if (s and s.inlet) else None), "t/h"])
            tbl_row(tbl, ["Isentropic Efficiency", spec.get("isentropic_efficiency"), ""])
            tbl_row(tbl, ["Mechanical Efficiency", spec.get("mechanical_efficiency"), ""])
            tbl_row(tbl, ["Generator Efficiency", spec.get("generator_efficiency"), ""])

            shaft = s.metadata.get("shaft_power_kw") if (s and s.metadata) else None
            gross = s.metadata.get("gross_electrical_kw") if (s and s.metadata) else None
            tbl_row(tbl, ["Shaft Power", _safe(shaft), "kW"])
            tbl_row(tbl, ["Gross Electrical", _safe(gross), "kW"])
            if gross:
                tbl_row(tbl, ["Gross Electrical", _safe(gross / 1000.0, ".2f"), "MW"])

            ext_levels = spec.get("extraction_levels", [])
            ext_flows = spec.get("extraction_flows_tph", [])
            if ext_levels:
                doc.add_paragraph()
                etbl = add_styled_table(doc, ["Extraction Port", "Header Level",
                                              "Design Flow (t/h)", "Solved Flow (t/h)"])
                for i, lvl in enumerate(ext_levels):
                    pname = f"extraction_{i+1}"
                    design_f = ext_flows[i] if i < len(ext_flows) else 0
                    solved_f = None
                    if s and s.extraction_states and pname in s.extraction_states:
                        solved_f = s.extraction_states[pname].mass_flow_tph
                    etbl.add_row().cells[0].text = ""  # spacer bug workaround
                    row = etbl.rows[-1]
                    for ci, v in enumerate([pname, lvl, _safe(design_f), _safe(solved_f)]):
                        row.cells[ci].text = str(v) if v is not None else "--"
    else:
        doc.add_paragraph("No turbines in the current network.")

    # ---- 5. PRDS ----
    doc.add_heading("5. PRDS Stations", level=1)
    prds_nodes = [n for n in nodes if n["type"] == "prds"]
    if prds_nodes:
        tbl = add_styled_table(doc, ["Name", "From", "To", "Max Cap (t/h)",
                                      "Solved In (t/h)", "Solved Out (t/h)", "Spray (t/h)"])
        for n in prds_nodes:
            spec = (n.get("metadata") or {}).get("spec", {})
            s = state_map.get(n["id"])
            m_in = s.inlet.mass_flow_tph if (s and s.inlet) else None
            m_out = s.outlet.mass_flow_tph if (s and s.outlet) else None
            spray = (m_out - m_in) if (m_in is not None and m_out is not None) else None
            tbl_row(tbl, [
                n["name"], spec.get("from_level"), spec.get("to_level"),
                spec.get("max_capacity_tph"),
                _safe(m_in), _safe(m_out), _safe(spray),
            ])
    else:
        doc.add_paragraph("No PRDS stations in the current network.")

    # ---- 6. CONSUMERS ----
    doc.add_heading("6. Steam Consumers", level=1)
    cons_nodes = [n for n in nodes if n["type"] == "consumer"]
    if cons_nodes:
        tbl = add_styled_table(doc, ["Name", "Header", "Demand (t/h)",
                                      "Return Frac", "Solved Flow (t/h)"])
        for n in cons_nodes:
            spec = (n.get("metadata") or {}).get("spec", {})
            s = state_map.get(n["id"])
            flow = _safe(s.inlet.mass_flow_tph if (s and s.inlet) else None)
            tbl_row(tbl, [
                n["name"], spec.get("header_level"), spec.get("demand_tph"),
                spec.get("return_fraction"), flow,
            ])
    else:
        doc.add_paragraph("No consumers in the current network.")

    # ---- 7. ALL COMPONENT THERMODYNAMIC STATES ----
    doc.add_heading("7. Thermodynamic States (All Components)", level=1)
    doc.add_paragraph(
        "Post-solve thermodynamic state at each component's inlet and outlet ports."
    )
    if result and result.component_states:
        tbl = add_styled_table(doc, ["Component", "Port", "P (bar)", "T (C)",
                                      "h (kJ/kg)", "s (kJ/kg.K)", "Flow (t/h)", "Phase"])
        for s in result.component_states:
            if s.inlet:
                si = s.inlet
                tbl_row(tbl, [s.name, "inlet",
                              _safe(si.pressure_bar), _safe(si.temperature_c),
                              _safe(si.enthalpy_kj_kg), _safe(si.entropy_kj_kgk, ".3f"),
                              _safe(si.mass_flow_tph),
                              si.phase.value if si.phase else "--"])
            if s.outlet:
                so = s.outlet
                tbl_row(tbl, [s.name, "outlet",
                              _safe(so.pressure_bar), _safe(so.temperature_c),
                              _safe(so.enthalpy_kj_kg), _safe(so.entropy_kj_kgk, ".3f"),
                              _safe(so.mass_flow_tph),
                              so.phase.value if so.phase else "--"])
            for pname, ex in (s.extraction_states or {}).items():
                tbl_row(tbl, [s.name, pname,
                              _safe(ex.pressure_bar), _safe(ex.temperature_c),
                              _safe(ex.enthalpy_kj_kg), _safe(ex.entropy_kj_kgk, ".3f"),
                              _safe(ex.mass_flow_tph),
                              ex.phase.value if ex.phase else "--"])
    else:
        doc.add_paragraph("No solved states available (run Solve first).")

    # ---- 8. CONNECTIONS ----
    doc.add_heading("8. Network Connections", level=1)
    if edges:
        tbl = add_styled_table(doc, ["From", "Port", "To", "Port", "P (bar)"])
        for e in edges:
            from_name = next((n["name"] for n in nodes if n["id"] == e["from"]), e["from"])
            to_name = next((n["name"] for n in nodes if n["id"] == e["to"]), e["to"])
            tbl_row(tbl, [from_name, e["from_port"], to_name, e["to_port"],
                          _safe(e.get("pressure_bar"))])
    else:
        doc.add_paragraph("No connections yet.")

    # ---- footer ----
    doc.add_page_break()
    doc.add_heading("Document Control", level=1)
    tbl = add_styled_table(doc, ["Field", "Value"])
    tbl_row(tbl, ["Generated", now.strftime("%Y-%m-%d %H:%M:%S")])
    tbl_row(tbl, ["System", "Steam Network Simulation Engine"])
    tbl_row(tbl, ["Solver Backend", "CoolProp / IAPWS-IF97"])
    tbl_row(tbl, ["Network Hash", f"{len(nodes)}n-{len(edges)}e"])
    tbl_row(tbl, ["Report Version", "Auto-generated (live)"])

    fpath = out / "Steam_Network_Report.docx"
    doc.save(str(fpath))
    log.info("Word report written: %s", fpath)
    return fpath
